"""Аудит всей выгрузки — воспроизводимые факты и локальные листы просмотра.

    python -m src.utils.audit

Исходники и метки не меняются. Результаты в artifacts/audit; привязка сторон
наследует гипотезу загрузчика, а контакты не заменяют экспертную переразметку.
"""

import hashlib
import importlib.metadata
import json
import platform
import zipfile
from collections import Counter
from datetime import datetime, timezone

import numpy as np
import openpyxl
import pandas as pd
import pydicom
from PIL import Image, ImageDraw
from sklearn.model_selection import GroupKFold

from .. import config as C
from . import data


OUT = C.ARTIFACTS / "audit"
TARGETS = ["spine_position", "spine_axis", "spine_objects", "hip_position", "hip_roi"]


def sha(payload):
    return hashlib.sha256(payload).hexdigest()


def counts(values):
    return dict(sorted(Counter(map(str, values)).items()))


def audit():
    OUT.mkdir(parents=True, exist_ok=True)
    summary = {"generated_utc": datetime.now(timezone.utc).isoformat(),
               "python": platform.python_version(),
               "versions": {p: importlib.metadata.version(p) for p in
                            ("numpy", "pandas", "pydicom", "Pillow", "scipy", "scikit-learn", "openpyxl")}}
    archives = []
    for filename, destination in (("НД_для_обучения.zip", C.TRAIN), ("Для теста.zip", C.TEST)):
        path = C.RAW / filename
        with zipfile.ZipFile(path) as archive:
            failures, names = [], []
            for member in archive.infolist():
                if member.is_dir():
                    continue
                try:
                    name = member.filename.encode("cp437").decode("cp866")
                except UnicodeError:
                    name = member.filename
                names.append(name)
                target = destination / name
                if not target.is_file() or sha(target.read_bytes()) != sha(archive.read(member)):
                    failures.append(name)
            archives.append(dict(name=filename, sha256=sha(path.read_bytes()),
                                 crc_error=archive.testzip(), files=len(names),
                                 extracted_mismatch=failures))
    summary["archives"] = archives

    book = pd.read_excel(C.LABELS_XLSX, header=None)
    flat, workbook_rows = [], []
    for _, r in book.iloc[2:].iterrows():
        if pd.isna(r[1]):
            continue
        case = f"S{int(r[0]):03}"
        workbook_rows.append(dict(case=case, study=r[1], excel_row=int(r.name)+1,
                                  comment="" if pd.isna(r[12]) else str(r[12])))
        for region, side, cols, total_col in (("spine", "", [2, 3, 4], 9),
                                              ("hip", "right", [5, 6], 10),
                                              ("hip", "left", [7, 8], 11)):
            vals = r[cols]
            known = vals.notna().all()
            row = dict(case=case, study=r[1], region=region, side=side,
                       flags_known=bool(known), partially_missing=bool(vals.notna().any() and not known),
                       flag_quality=int(vals.max()) if known else np.nan,
                       expert_quality=r[total_col], excel_row=int(r.name)+1,
                       comment="" if pd.isna(r[12]) else str(r[12]))
            row.update(dict(zip(TARGETS[:3] if region == "spine" else TARGETS[3:], vals)))
            flat.append(row)
    labels = pd.DataFrame(flat)
    labels.to_csv(OUT / "labels.csv", index=False)
    study_cases = {r["study"]: r["case"] for r in workbook_rows}
    mismatch = labels[labels.flag_quality.notna() & labels.expert_quality.notna() &
                      (labels.flag_quality != labels.expert_quality)]
    summary["workbook"] = dict(sha256=sha(C.LABELS_XLSX.read_bytes()),
        studies=len(workbook_rows), duplicate_study_rows=int(pd.Series([r["study"] for r in workbook_rows]).duplicated().sum()),
        flag_values=counts(book.iloc[2:, 2:9].to_numpy().ravel()),
        partially_missing=int(labels.partially_missing.sum()),
        totals_disagree=mismatch.fillna("").to_dict("records"),
        comments=[r for r in workbook_rows if r["comment"]])
    workbook = openpyxl.load_workbook(C.LABELS_XLSX, data_only=False)
    cached = openpyxl.load_workbook(C.LABELS_XLSX, data_only=True)
    summary["workbook"]["sheets"] = [dict(name=s.title, rows=s.max_row, columns=s.max_column,
        hidden=s.sheet_state, hidden_rows=[i for i, r in s.row_dimensions.items() if r.hidden],
        hidden_columns=[i for i, c in s.column_dimensions.items() if c.hidden]) for s in workbook]
    summary["workbook"]["formulas"] = [dict(cell=c.coordinate, formula=c.value,
        cached=cached[s.title][c.coordinate].value) for s in workbook for row in s for c in row if c.data_type == "f"]

    rows, arrays, errors = [], {}, []
    dicom_tags = ["Modality", "Manufacturer", "ManufacturerModelName", "TransferSyntaxUID",
                  "BitsAllocated", "BitsStored", "HighBit", "PixelRepresentation", "SamplesPerPixel",
                  "PhotometricInterpretation", "NumberOfFrames", "PixelSpacing", "ImagerPixelSpacing",
                  "BodyPartExamined", "ViewPosition", "Laterality", "ImageLaterality", "PatientOrientation",
                  "RescaleSlope", "RescaleIntercept", "WindowCenter", "WindowWidth", "BurnedInAnnotation",
                  "SOPClassUID", "ImageType", "StudyDescription", "SeriesDescription", "ProtocolName"]
    tags = {k: [] for k in dicom_tags}
    private_tags, anonymization = Counter(), Counter()
    for split, directory in (("train", C.STUDIES), ("test", C.TEST)):
        for path in sorted(directory.rglob("*.dcm")):
            rel = str(path.relative_to(C.DATA))
            try:
                d = pydicom.dcmread(path)
                a = d.pixel_array
                assert a.ndim == 2, a.shape
                digest = sha(str(a.shape).encode() + str(a.dtype).encode() + a.tobytes())
                study = path.relative_to(directory).parts[0] if split == "train" else str(d.get("StudyInstanceUID", ""))
                region = "spine" if a.shape[1] == C.SPINE_COLUMNS else "hip"
                score = data.side_score(a)
                values = np.unique(a)
                row = dict(split=split, path_to_study=rel, study=study,
                           case=study_cases.get(study, "TEST"), region=region,
                           side="" if region == "spine" else ("left" if score > 0 else "right"),
                           pixel_sha256=digest, file_sha256=sha(path.read_bytes()),
                           study_uid=str(d.get("StudyInstanceUID", "")),
                           series_uid=str(d.get("SeriesInstanceUID", "")), image_uid=str(d.get("SOPInstanceUID", "")),
                           uid_valid=all(getattr(d.get(k, ""), "is_valid", False) for k in
                                         ("StudyInstanceUID", "SeriesInstanceUID", "SOPInstanceUID")),
                           rows=a.shape[0], columns=a.shape[1], dtype=str(a.dtype),
                           minimum=int(a.min()), maximum=int(a.max()), mean=float(a.mean()), std=float(a.std()),
                           zero_fraction=float((a == 0).mean()), saturation_fraction=float((a >= 250).mean()),
                           side_score=score,
                           constant=bool(len(values) == 1))
                for frac in (0.2, 0.5):
                    top = a[:int(len(a)*frac)].astype(float)
                    row[f"side_{frac}"] = float(top[:, :a.shape[1]//2].mean() - top[:, a.shape[1]//2:].mean())
                rows.append(row)
                arrays.setdefault(digest, a)
                for key in tags:
                    tags[key].append(str(d.file_meta.get(key, d.get(key, "<missing>"))))
                for elem in d.iterall():
                    if elem.tag.is_private:
                        private_tags[str(elem.tag)] += 1
                    if elem.VR in ("PN", "DA", "TM"):
                        value = str(elem.value)
                        anonymization[f"{elem.VR}:" + ("Anonymized" if value == "Anonymized" else "empty" if not value else "other")] += 1
            except Exception as exc:
                errors.append(dict(path=rel, error=f"{type(exc).__name__}: {exc}"))
    files = pd.DataFrame(rows)
    files.to_csv(OUT / "files.csv", index=False)
    summary["dicom"] = dict(files=len(files), errors=errors, tags={k: counts(v) for k, v in tags.items()},
        private_tags=dict(private_tags), anonymization=dict(anonymization), invalid_uid_files=int((~files.uid_valid).sum()),
        constant_files=int(files.constant.sum()),
        uid_distinct={k: int(files[k].nunique()) for k in ("study_uid", "series_uid", "image_uid")})
    train = files[files.split == "train"]
    unique = files.drop_duplicates(["split", "pixel_sha256"]).copy()
    unique = unique.merge(labels, on=["case", "study", "region", "side"], how="left", validate="many_to_one")
    unique["image_id"] = unique.case + "_" + unique.region + "_" + unique.side
    unique.to_csv(OUT / "images.csv", index=False)
    differing_tags = Counter()
    for _, frame in train.groupby("pixel_sha256"):
        if len(frame) < 2:
            continue
        copies = [pydicom.dcmread(C.DATA / p, stop_before_pixels=True) for p in frame.path_to_study]
        for tag in set().union(*(set(d.keys()) for d in copies)):
            if len({str(d.get(tag)) for d in copies}) > 1:
                differing_tags[pydicom.datadict.keyword_for_tag(tag) or str(tag)] += 1
    summary["duplicates"] = dict(train_files=len(train), train_pixels=int(train.pixel_sha256.nunique()),
        train_bytes=int(train.file_sha256.nunique()), test_files=int((files.split == "test").sum()),
        test_pixels=int(files[files.split == "test"].pixel_sha256.nunique()),
        cross_study_pixel_groups=int((train.groupby("pixel_sha256").study.nunique() > 1).sum()),
        cross_split_pixel_groups=int((files.groupby("pixel_sha256").split.nunique() > 1).sum()),
        uid_conflicts=int((files.groupby("image_uid").pixel_sha256.nunique() > 1).sum()),
        differing_tags=dict(differing_tags),
        duplicate_groups=int((train.groupby("pixel_sha256").size() > 1).sum()),
        cross_series_pixel_groups=int((train.groupby("pixel_sha256").series_uid.nunique() > 1).sum()))
    summary["study_mapping"] = dict(folder_equals_uid=int((train.study == train.study_uid).sum()),
        max_uids_per_folder=int(train.groupby("study").study_uid.nunique().max()),
        max_folders_per_uid=int(train.groupby("study_uid").study.nunique().max()))
    matched = labels.merge(unique[unique.split == "train"][["study", "region", "side", "image_id"]],
                           on=["study", "region", "side"], how="left", validate="one_to_one")
    missing_images = matched[matched.flags_known & matched.image_id.isna()]
    missing_labels = unique[(unique.split == "train") & unique.flag_quality.isna()]
    summary["alignment"] = dict(images_without_flags=missing_labels[["image_id", "study", "path_to_study"]].to_dict("records"),
        flags_without_images=missing_images.fillna("").to_dict("records"))
    u = unique[unique.split == "train"].copy()
    summary["composition"] = dict(images=len(u), studies=int(u.study.nunique()),
        shapes=counts(zip(u.rows, u["columns"])), study_sizes=counts(u.groupby("study").size()),
        region_sides=counts(u.region + "_" + u.side), labeled=int(u.flag_quality.notna().sum()),
        quality_flags=int(u.flag_quality.sum()), quality_expert=int(u.expert_quality.sum()),
        positive_studies_flags=int(u.groupby("study").flag_quality.max().sum()),
        positive_studies_expert=int(u.groupby("study").expert_quality.max().sum()),
        targets={t: dict(positive_images=int(u[t].sum()), known_images=int(u[t].notna().sum()),
                         positive_studies=int(u.groupby("study")[t].max().sum())) for t in TARGETS})
    summary["intensity"] = {region: frame[["rows", "columns", "mean", "std", "zero_fraction", "saturation_fraction"]]
                            .describe().to_dict() for region, frame in u.groupby("region")}
    summary["side"] = dict(min_abs=float(u[u.region == "hip"].side_score.abs().min()),
        unstable=u[(u.region == "hip") & ((np.sign(u.side_score) != np.sign(u["side_0.2"])) |
                                          (np.sign(u.side_score) != np.sign(u["side_0.5"])))].image_id.tolist())
    paired = labels[labels.region == "hip"].pivot(index="study", columns="side", values=TARGETS[3:]).dropna()
    summary["side"]["fully_labeled_pairs"] = len(paired)
    summary["side"]["different_labels"] = int((paired.xs("left", axis=1, level=1) !=
                                                paired.xs("right", axis=1, level=1)).any(axis=1).sum())
    hip_counts = u[u.region == "hip"].groupby("study").side.nunique()
    summary["side"]["opposite_sign_pairs"] = int((hip_counts == 2).sum())
    summary["cooccurrence"] = u[TARGETS].fillna(0).astype(int).T.dot(u[TARGETS].fillna(0).astype(int)).to_dict()
    conflicts = set(mismatch.case)
    u["excluded_primary"] = u.flag_quality.isna() | ((u.region == "spine") & u.case.isin(conflicts))
    u.to_csv(OUT / "training_candidates.csv", index=False)
    summary["primary_candidates"] = dict(images=int((~u.excluded_primary).sum()),
        studies=int(u[~u.excluded_primary].study.nunique()),
        targets=u[~u.excluded_primary][TARGETS].sum().astype(int).to_dict())

    # Nearest neighbours are review candidates, not proof of patient identity.
    thumb = np.stack([np.asarray(Image.fromarray(arrays[h]).resize((48, 64)), dtype=float).ravel()
                      for h in unique.pixel_sha256])
    thumb -= thumb.mean(axis=1, keepdims=True)
    thumb /= np.maximum(np.linalg.norm(thumb, axis=1, keepdims=True), 1e-12)
    corr = thumb @ thumb.T
    # ponytail: O(n²) для 255 кадров; индекс соседей нужен при росте набора.
    pairs = []
    for i in range(len(unique)):
        for j in range(i+1, len(unique)):
            x, y = unique.iloc[i], unique.iloc[j]
            if x.study != y.study and x.region == y.region:
                pairs.append(dict(image_a=x.image_id, image_b=y.image_id,
                                  path_a=x.path_to_study, path_b=y.path_to_study,
                                  correlation=float(corr[i, j]), cross_split=x.split != y.split))
    nearest = pd.DataFrame(pairs).sort_values("correlation", ascending=False)
    nearest.head(50).to_csv(OUT / "nearest.csv", index=False)
    summary["nearest"] = nearest.head(10).to_dict("records")
    summary["nearest_cross_split"] = nearest[nearest.cross_split].head(5).to_dict("records")

    summary["folds"] = {}
    for variant, eligible, k in (("raw_3", u, 3), ("raw_5", u, 5),
                                 ("primary_3", u[~u.excluded_primary], 3)):
        groups = eligible.groupby("study")[TARGETS].max().fillna(0).astype(int)
        default = []
        labeled = eligible[eligible.flag_quality.notna()]
        for _, ix in GroupKFold(k).split(labeled, groups=labeled.study):
            default.append(labeled.iloc[ix].groupby("study")[TARGETS].max().fillna(0).sum().astype(int).tolist())
        # Select by label balance only; never by model scores. Provisional until labels are resolved.
        rng, best = np.random.default_rng(42), None
        for _ in range(10000):
            assignment = np.empty(len(groups), dtype=int)
            assignment[rng.permutation(len(groups))] = np.arange(len(groups)) % k
            totals = np.stack([groups.to_numpy()[assignment == f].sum(axis=0) for f in range(k)])
            score = float(((totals - groups.sum().to_numpy()/k)**2 / np.maximum(groups.sum().to_numpy(), 1)).sum())
            score += 100 * int((totals == 0).sum())
            if best is None or score < best[0]:
                best = (score, assignment.copy(), totals)
        pd.DataFrame(dict(study=groups.index, fold=best[1])).to_csv(OUT / f"candidate_folds_{variant}.csv", index=False)
        summary["folds"][variant] = dict(targets=TARGETS, default_positive_studies=default,
                                        candidate_positive_studies=best[2].tolist(),
                                        candidate_study_sizes=counts(best[1]))

    for page, start in enumerate(range(0, len(unique), 16), 1):
        canvas = Image.new("RGB", (4*300, 4*440), "#202020")
        draw = ImageDraw.Draw(canvas)
        for slot, (_, row) in enumerate(unique.iloc[start:start+16].iterrows()):
            x, y = slot % 4 * 300, slot // 4 * 440
            a = Image.fromarray(arrays[row.pixel_sha256]).convert("RGB")
            a.thumbnail((292, 390))
            canvas.paste(a, (x+(300-a.width)//2, y+45))
            flags = "unknown" if pd.isna(row.flag_quality) else ",".join(
                t.split("_")[-1] for t in TARGETS if row.get(t) == 1) or "none"
            draw.text((x+5, y+5), row.image_id, fill="white")
            draw.text((x+5, y+22), f"flags={flags} total={row.expert_quality}", fill="#ffcc66")
        canvas.save(OUT / f"contact_{page:02}.jpg", quality=92)

    current = data.table().set_index("path_to_study")
    independent = u.set_index("path_to_study").flag_quality.reindex(current.index)
    assert np.array_equal(current.quality_class.to_numpy(), independent.to_numpy(), equal_nan=True)
    assert len(files) == sum(len(list(p.rglob("*.dcm"))) for p in (C.STUDIES, C.TEST)) - len(errors)
    assert int(u[TARGETS].eq(1).any(axis=1).sum()) == int(u.flag_quality.sum())
    assert sum(summary["composition"]["shapes"].values()) == len(u)
    assert not labels.partially_missing.any(), "Partial labels require target-specific handling"
    assert set(book.iloc[2:, 2:9].stack().dropna().unique()) <= {0, 1}
    assert all(not a["extracted_mismatch"] and a["crc_error"] is None for a in archives)
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False))
    print(json.dumps({k: summary[k] for k in ("archives", "duplicates", "composition", "alignment", "side", "folds")},
                     ensure_ascii=False, indent=2))
    print(f"Workbook summary conflicts: {len(mismatch)}; outputs: {OUT}")


if __name__ == "__main__":
    audit()
